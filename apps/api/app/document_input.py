"""Bounded, page-aware input preparation. Never consults benchmark labels."""

import base64
import csv
import io
import json
import os
import re
import subprocess
import tempfile
import time
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image
from pypdf import PdfReader

PIPELINE_VERSION = "local-2.1"


@dataclass
class DocumentInput:
    text: str
    images: list[str] = field(default_factory=list)
    image_pages: list[int] = field(default_factory=list)
    pages: dict[int, str] = field(default_factory=dict)
    metadata: dict = field(default_factory=dict)


def image_data(path):
    with Image.open(path) as source:
        if source.width * source.height > 40_000_000:
            raise ValueError("Image exceeds 40 megapixels.")
        source = source.convert("RGB")
        source.thumbnail((1600, 1600))
        stream = io.BytesIO()
        source.save(stream, format="PNG")
        return base64.b64encode(stream.getvalue()).decode()


def prepare_document(path: Path) -> DocumentInput:
    from .providers import parser_for, TesseractOCRProvider

    started = time.monotonic()
    result = DocumentInput("")
    ocr_seconds = 0.0
    limit = max(1, min(4, int(os.getenv("AI_MAX_PAGES", "3"))))
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(path)
        if reader.is_encrypted or len(reader.pages) > 50:
            raise ValueError("Use an unlocked PDF with at most 50 pages.")
        pages = {i + 1: p.extract_text(extraction_mode="layout") or "" for i, p in enumerate(reader.pages)}
        selection_text = dict(pages)
        if len(pages) > limit and sum(len(p.strip()) < 80 for p in pages.values()) > limit:
            selection_started = time.monotonic()
            with tempfile.TemporaryDirectory() as temporary:
                prefix = Path(temporary) / "thumb"
                subprocess.run(
                    ["pdftoppm", "-png", "-scale-to", "700", str(path), str(prefix)],
                    check=True,
                    capture_output=True,
                    timeout=90,
                )
                for thumbnail in sorted(Path(temporary).glob("thumb-*.png")):
                    number = int(thumbnail.stem.rsplit("-", 1)[1])
                    if time.monotonic() - selection_started > 60:
                        result.metadata["page_selection_incomplete"] = True
                        break
                    if len(pages[number].strip()) < 80:
                        selection_text[number] = TesseractOCRProvider().extract(thumbnail)
            selection_seconds = time.monotonic() - selection_started
            ocr_seconds += selection_seconds
            result.metadata["page_selection_ocr_seconds"] = round(selection_seconds, 4)

        # Always retain identity on page one; rank remaining pages by commercial signals.
        def score(number):
            return len(
                re.findall(
                    r"quote|quotation|unit price|quantity|sku|moq|delivery|subtotal|total|pricing",
                    selection_text[number],
                    re.I,
                )
            )

        selected = (
            sorted({1, *sorted(range(2, len(pages) + 1), key=lambda n: (-score(n), n))[: limit - 1]}) if pages else []
        )
        result.metadata.update(
            page_count=len(pages), selected_pages=selected, omitted_pages=sorted(set(pages) - set(selected))
        )
        for number in selected:
            content = pages[number]
            if len(content.strip()) < 80:
                with tempfile.TemporaryDirectory() as temporary:
                    prefix = Path(temporary) / "page"
                    subprocess.run(
                        [
                            "pdftoppm",
                            "-f",
                            str(number),
                            "-l",
                            str(number),
                            "-singlefile",
                            "-png",
                            "-scale-to",
                            "1600",
                            str(path),
                            str(prefix),
                        ],
                        check=True,
                        capture_output=True,
                        timeout=45,
                    )
                    image_path = prefix.with_suffix(".png")
                    t = time.monotonic()
                    content = TesseractOCRProvider().extract(image_path)
                    ocr_seconds += time.monotonic() - t
                    result.images.append(image_data(image_path))
                    result.image_pages.append(number)
            result.pages[number] = content
        result.text = "\n".join(f"[Page {n}]\n{p}" for n, p in result.pages.items())
    elif suffix in (".png", ".jpg", ".jpeg"):
        result.images = [image_data(path)]
        result.image_pages = [1]
        t = time.monotonic()
        try:
            content = TesseractOCRProvider().extract(path)
        except (RuntimeError, subprocess.TimeoutExpired):
            content = ""
        ocr_seconds = time.monotonic() - t
        result.pages = {1: content}
        result.text = "[Page 1]\n" + content
    elif suffix in (".xlsx", ".csv"):
        # Existing parsers enforce ZIP expansion and row limits first.
        parser_for(path.name).parse(path)
        tables = []
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                tables = [(sheet.title, list(sheet.iter_rows(values_only=True))) for sheet in workbook]
            finally:
                workbook.close()
        else:
            with path.open(encoding="utf-8-sig", newline="") as stream:
                tables = [("CSV", list(csv.reader(stream)))]
        chunks = []
        for name, rows in tables:
            headers = None
            for index, raw in enumerate(rows):
                row = list(raw)
                while row and row[-1] is None:
                    row.pop()
                labels = [str(value or "").strip().lower() for value in row]
                if "sku" in labels and ("qty" in labels or "quantity" in labels) and "unit price" in labels:
                    headers = list(map(str, row))
                record = (
                    dict(zip(headers, row))
                    if headers and len(headers) == len(row) and labels != [h.lower() for h in headers]
                    else row
                )
                chunks.append(
                    json.dumps({"sheet": name, "row": index + 1, "cells": record}, ensure_ascii=False, default=str)
                )
        result.text = "Structured spreadsheet rows; empty cells are missing values.\n" + "\n".join(chunks)
    else:
        result.text = parser_for(path.name).parse(path)
    # Reject oversize input rather than silently extract a truncated quotation.
    if len(result.text) > 18000:
        raise ValueError(
            "Selected quotation text exceeds the local context budget. Split the quote into smaller documents."
        )
    result.metadata.update(
        pipeline_version=PIPELINE_VERSION,
        parsing_seconds=round(time.monotonic() - started - ocr_seconds, 4),
        ocr_seconds=round(ocr_seconds, 4),
        input_characters=len(result.text),
        image_count=len(result.images),
        image_pages=result.image_pages,
        parser_quality="low" if len(result.text.strip()) < 80 else "normal",
    )
    return result
