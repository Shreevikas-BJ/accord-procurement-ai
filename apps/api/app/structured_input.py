"""Format-specific, row-preserving candidates for local quote extraction.

Candidates contain source structure only. They never use benchmark labels or
catalog/history data, and they are safe to serialize into the untrusted model
envelope and extraction diagnostics.
"""

from __future__ import annotations

import csv
import re
from pathlib import Path


HEADER_ALIASES = {
    "supplier_sku": (
        "sku",
        "supplier sku",
        "supplier part number",
        "supplier part no",
        "vendor part number",
        "vendor part no",
        "part number",
        "part no",
        "item number",
        "item no",
        "item code",
        "material",
        "stock no",
    ),
    "description": ("description", "item description", "product", "product description"),
    "quantity": ("qty", "oty", "quantity", "quoted quantity", "order qty", "order quantity"),
    "uom": ("uom", "yom", "vom", "unit of measure", "unit"),
    "unit_price": (
        "unit price",
        "price per unit",
        "price each",
        "each price",
        "unit cost",
        "net ea",
        "net ea.",
    ),
    "stated_line_total": ("line total", "extended", "extended price", "extension", "amount"),
    "moq": ("moq", "moqg", "minimum order quantity", "minimum qty", "min qty"),
    "lead_time_days": ("lead time", "lead time days", "lead", "availability"),
    "delivery_date": ("delivery", "delivery date", "ship date", "promise date"),
}


def normalized_label(value) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").casefold()).strip()


HEADER_LOOKUP = {
    normalized_label(alias): field
    for field, aliases in HEADER_ALIASES.items()
    for alias in (field, *aliases)
}


def semantic_header(value) -> str | None:
    return HEADER_LOOKUP.get(normalized_label(value))


def column_name(index: int) -> str:
    result = ""
    while index >= 0:
        result = chr(index % 26 + 65) + result
        index = index // 26 - 1
    return result


def _clean_row(values) -> list:
    row = list(values)
    while row and row[-1] is None:
        row.pop()
    return row


def spreadsheet_candidates(path: Path) -> tuple[list[dict], dict]:
    """Preserve sheet, physical row, raw cells, semantic headers and cell refs."""
    suffix = path.suffix.casefold()
    candidates: list[dict] = []
    merged_by_sheet: dict[str, list[str]] = {}
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=False, data_only=True)
        try:
            tables = [(sheet.title, [_clean_row(r) for r in sheet.iter_rows(values_only=True)]) for sheet in workbook]
            merged_by_sheet = {
                sheet.title: [str(cell_range) for cell_range in sheet.merged_cells.ranges]
                for sheet in workbook
                if sheet.merged_cells.ranges
            }
        finally:
            workbook.close()
    else:
        with path.open(encoding="utf-8-sig", newline="") as stream:
            tables = [("CSV", [_clean_row(r) for r in csv.reader(stream)])]

    for sheet, rows in tables:
        headers: list[str | None] | None = None
        header_values: list[str] | None = None
        table_id = 0
        for index, row in enumerate(rows, 1):
            semantic = [semantic_header(value) for value in row]
            recognized = sum(value is not None for value in semantic)
            is_header = recognized >= 3 and any(value in semantic for value in ("supplier_sku", "description"))
            if is_header:
                table_id += 1
                headers = semantic
                header_values = [str(value or "") for value in row]
            field_cells = {}
            if headers and not is_header:
                for column, key in enumerate(headers):
                    if key and column < len(row):
                        field_cells[key] = {
                            "value": row[column],
                            "cell": f"{sheet}!{column_name(column)}{index}",
                        }
            candidates.append(
                {
                    "source_type": "xlsx" if suffix == ".xlsx" else "csv",
                    "sheet": sheet,
                    "table": table_id or None,
                    "row": index,
                    "headers": header_values if headers else None,
                    "semantic_headers": headers,
                    "cells": row,
                    "cell_refs": [f"{sheet}!{column_name(i)}{index}" for i in range(len(row))],
                    "field_cells": field_cells,
                    "raw_text": " | ".join("" if value is None else str(value) for value in row),
                    "layout_confidence": "strong" if field_cells else "unmapped",
                }
            )
    return candidates, {"merged_cells": merged_by_sheet}


def word_rows(words: list[dict], *, source_type: str, page: int) -> list[dict]:
    """Group pdf/OCR words into geometric lines and map rows beneath a header."""
    groups: dict[tuple, list[dict]] = {}
    for word in words:
        text = str(word.get("text") or "").strip()
        if not text:
            continue
        key = word.get("line_key")
        if key is None:
            # pdfplumber words do not expose OCR line ids; y bucketing is stable
            # enough to preserve physical rows without asserting semantics.
            key = (round(float(word.get("top", 0)) / 4),)
        groups.setdefault(tuple(key) if isinstance(key, (list, tuple)) else (key,), []).append(word)
    rows = []
    for row_index, group in enumerate(
        sorted(groups.values(), key=lambda values: (min(float(w.get("top", 0)) for w in values), min(float(w.get("x0", 0)) for w in values))),
        1,
    ):
        ordered = sorted(group, key=lambda word: float(word.get("x0", 0)))
        text = " ".join(str(word["text"]).strip() for word in ordered)
        rows.append(
            {
                "source_type": source_type,
                "page": page,
                "row": row_index,
                "cells": [word["text"] for word in ordered],
                "words": [
                    {
                        "text": word["text"],
                        "bbox": [
                            round(float(word.get("x0", 0)), 2),
                            round(float(word.get("top", 0)), 2),
                            round(float(word.get("x1", 0)), 2),
                            round(float(word.get("bottom", 0)), 2),
                        ],
                    }
                    for word in ordered
                ],
                "bbox": [
                    round(min(float(word.get("x0", 0)) for word in ordered), 2),
                    round(min(float(word.get("top", 0)) for word in ordered), 2),
                    round(max(float(word.get("x1", 0)) for word in ordered), 2),
                    round(max(float(word.get("bottom", 0)) for word in ordered), 2),
                ],
                "raw_text": text,
                "layout_confidence": "unmapped",
            }
        )
    _map_geometric_tables(rows)
    return rows


def _header_positions(row: dict) -> list[tuple[str, float, str]]:
    words = row.get("words") or []
    result: list[tuple[str, float, str]] = []
    for start in range(len(words)):
        for length in (3, 2, 1):
            phrase = " ".join(str(w["text"]) for w in words[start : start + length])
            key = semantic_header(phrase)
            if key:
                x = float(words[start]["bbox"][0])
                if key not in {entry[0] for entry in result}:
                    result.append((key, x, phrase))
                break
    return sorted(result, key=lambda entry: entry[1])


def _map_geometric_tables(rows: list[dict]) -> None:
    header: list[tuple[str, float, str]] | None = None
    pipe_headers: list[str | None] | None = None
    table_id = 0
    for row in rows:
        pipe_parts = [part.strip() for part in row.get("raw_text", "").split("|")]
        pipe_semantic = [semantic_header(part) for part in pipe_parts]
        if len(pipe_parts) >= 3 and sum(key is not None for key in pipe_semantic) >= 3:
            table_id += 1
            pipe_headers = pipe_semantic
            header = None
            row["table"] = table_id
            row["headers"] = pipe_parts
            row["semantic_headers"] = pipe_semantic
            row["layout_confidence"] = "header"
            continue
        if pipe_headers and len(pipe_parts) == len(pipe_headers) and len(pipe_parts) >= 3:
            row["table"] = table_id
            row["field_cells"] = {
                key: {"value": pipe_parts[index] or None, "bbox": row.get("bbox")}
                for index, key in enumerate(pipe_headers)
                if key
            }
            row["layout_confidence"] = "strong"
            continue
        positions = _header_positions(row)
        if len(positions) >= 3 and any(key in {"supplier_sku", "description"} for key, _, _ in positions):
            table_id += 1
            header = positions
            pipe_headers = None
            row["table"] = table_id
            row["headers"] = [phrase for _, _, phrase in positions]
            row["semantic_headers"] = [key for key, _, _ in positions]
            row["layout_confidence"] = "header"
            continue
        if not header or not row.get("words"):
            continue
        if ":" in row.get("raw_text", "") and len(row.get("words") or []) < len(header):
            continue
        bounds = [-float("inf"), *[header[i + 1][1] - 2 for i in range(len(header) - 1)], float("inf")]
        assigned: dict[str, list[dict]] = {key: [] for key, _, _ in header}
        for word in row["words"]:
            center = (float(word["bbox"][0]) + float(word["bbox"][2])) / 2
            for index, (key, _, _) in enumerate(header):
                if bounds[index] <= center < bounds[index + 1]:
                    assigned[key].append(word)
                    break
        field_cells = {}
        for key, values in assigned.items():
            if values:
                field_cells[key] = {
                    "value": " ".join(str(value["text"]) for value in values),
                    "bbox": [
                        min(value["bbox"][0] for value in values),
                        min(value["bbox"][1] for value in values),
                        max(value["bbox"][2] for value in values),
                        max(value["bbox"][3] for value in values),
                    ],
                }
            else:
                field_cells[key] = {"value": None, "bbox": None}
        nonempty = sum(cell["value"] not in (None, "") for cell in field_cells.values())
        if nonempty >= 2:
            row["table"] = table_id
            row["field_cells"] = field_cells
            row["layout_confidence"] = "strong" if nonempty >= 4 else "weak"


def pdf_candidates(path: Path, selected_pages: list[int]) -> list[dict]:
    import pdfplumber

    candidates: list[dict] = []
    with pdfplumber.open(path) as pdf:
        for page_number in selected_pages:
            page = pdf.pages[page_number - 1]
            tables = page.extract_tables() or []
            page_has_mapped_table = False
            for table_index, table in enumerate(tables, 1):
                semantic_headers = None
                raw_headers = None
                for row_index, cells in enumerate(table, 1):
                    semantic = [semantic_header(cell) for cell in cells]
                    if sum(key is not None for key in semantic) >= 3:
                        semantic_headers = semantic
                        raw_headers = [str(cell or "") for cell in cells]
                    field_cells = {}
                    if semantic_headers and row_index > 1 and semantic != semantic_headers:
                        for column, key in enumerate(semantic_headers):
                            if key and column < len(cells):
                                field_cells[key] = {"value": cells[column], "bbox": None}
                    page_has_mapped_table = page_has_mapped_table or bool(field_cells)
                    candidates.append(
                        {
                            "source_type": "pdf_table",
                            "page": page_number,
                            "table": table_index,
                            "row": row_index,
                            "cells": [cell for cell in cells],
                            "headers": raw_headers,
                            "semantic_headers": semantic_headers,
                            "field_cells": field_cells,
                            "raw_text": " | ".join(str(cell or "") for cell in cells),
                            "layout_confidence": "table" if field_cells else "header" if semantic_headers == semantic else "unmapped",
                        }
                    )
            # pdfplumber tables and geometric word rows describe the same physical
            # rows. Prefer the explicit table when it mapped successfully so a
            # repeated lot is retained once rather than duplicated by two parsers.
            if not page_has_mapped_table:
                words = page.extract_words(keep_blank_chars=False, use_text_flow=False) or []
                candidates.extend(word_rows(words, source_type="pdf_words", page=page_number))
    return candidates


def ocr_candidates(path: Path, page: int = 1) -> list[dict]:
    import pytesseract
    from PIL import Image

    with Image.open(path) as image:
        data = pytesseract.image_to_data(
            image,
            config="--psm 6 -c preserve_interword_spaces=1",
            output_type=pytesseract.Output.DICT,
            timeout=45,
        )
    words = []
    for index, text in enumerate(data["text"]):
        if not str(text).strip() or float(data["conf"][index]) < 0:
            continue
        left, top = data["left"][index], data["top"][index]
        width, height = data["width"][index], data["height"][index]
        words.append(
            {
                "text": text,
                "x0": left,
                "top": top,
                "x1": left + width,
                "bottom": top + height,
                "line_key": (data["block_num"][index], data["par_num"][index], data["line_num"][index]),
            }
        )
    return word_rows(words, source_type="ocr_layout", page=page)


def compact_candidates(rows: list[dict], limit: int = 250) -> list[dict]:
    """Bound model context; full candidates remain available to evidence code."""
    compact = []
    model_rows = [
        row
        for row in rows
        if row.get("source_type") not in {"ocr_layout", "pdf_words"}
        or row.get("field_cells")
        or row.get("layout_confidence") == "header"
    ]
    for row in model_rows[:limit]:
        keys = [
            "source_type",
            "page",
            "sheet",
            "table",
            "row",
            "headers",
            "semantic_headers",
            "field_cells",
            "bbox",
            "raw_text",
            "layout_confidence",
        ]
        if not row.get("field_cells"):
            keys.extend(["cells", "cell_refs"])
        compact.append(
            {
                key: row[key]
                for key in keys
                if row.get(key) is not None
            }
        )
    return compact
